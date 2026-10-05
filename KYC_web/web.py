import os
import io
import streamlit as st
from docx import Document

st.set_page_config(page_title="نظام تعديل نماذج KYC", page_icon="📑", layout="centered")

st.title("📑 نظام معالجة وتحديث نماذج KYC")
st.write("قم باختيار النموذج وإدخال البيانات المطلوبة ثم تحميل المستند الناتج مباشرة.")

TEMPLATES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates")

def get_templates():
    if not os.path.exists(TEMPLATES_DIR):
        os.makedirs(TEMPLATES_DIR)
    return [f for f in os.listdir(TEMPLATES_DIR) if f.endswith(".docx") and not f.startswith("~$")]

template_files = get_templates()

if not template_files:
    st.error("لم يتم العثور على أي ملفات Word داخل مجلد `templates`!")
else:
    selected_template = st.selectbox("اختر النموذج المطلوب معالجته:", template_files)

    st.markdown("---")
    st.subheader("إدخال البيانات المتغيرة")

    col1, col2 = st.columns(2)

    with col1:
        kyc_subject = st.text_input("KYC Subject / Main Dealer")
        client_name_header = st.text_input("Client Name (Header)")
        client_id = st.text_input("Client ID")
        mcc = st.text_input("MCC")

    with col2:
        region = st.text_input("Region")
        client_name_cell = st.text_input("Client Name (Table)")
        nature_of_activity = st.text_input("Nature of Activity")

    data = {
        "KYC_SUBJECT": kyc_subject,
        "REGION": region,
        "CLIENT_NAME_HEADER": client_name_header,
        "CLIENT_NAME_CELL": client_name_cell,
        "CLIENT_ID": client_id,
        "NATURE_OF_ACTIVITY": nature_of_activity,
        "MCC": mcc
    }

    st.markdown("---")

    def replace_placeholders(doc, data):
        replacements = {f"{{{{{key}}}}}" : val for key, val in data.items()}
        for p in doc.paragraphs:
            for placeholder, value in replacements.items():
                if placeholder in p.text:
                    p.text = p.text.replace(placeholder, value)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for placeholder, value in replacements.items():
                            if placeholder in p.text:
                                p.text = p.text.replace(placeholder, value)

    if st.button("🚀 معالجة وتجهيز المستند", type="primary"):
        template_path = os.path.join(TEMPLATES_DIR, selected_template)
        try:
            doc = Document(template_path)
            replace_placeholders(doc, data)

            bio = io.BytesIO()
            doc.save(bio)
            bio.seek(0)

            st.success("تم تحديث المستند بنجاح!")
            st.download_button(
                label="📥 تنزيل المستند المعدل",
                data=bio,
                file_name=f"Updated_{selected_template}",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
        except Exception as e:
            st.error(f"حدث خطأ أثناء معالجة المستند: {e}")