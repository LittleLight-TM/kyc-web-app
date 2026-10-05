import os
import io
import subprocess
import streamlit as st
from docx import Document

# إعداد الصفحة
st.set_page_config(page_title="نظام تعديل نماذج KYC", page_icon="doc", layout="centered")

st.title("نظام معالجة وتحديث نماذج KYC")
st.write("قم باختيار النموذج وإدخال البيانات المطلوبة ثم تحميل المستند الناتج بصيغة PDF مباشرة.")

# المجلد المخصص للقوالب
TEMPLATES_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates")

def get_templates():
    if not os.path.exists(TEMPLATES_DIR):
        os.makedirs(TEMPLATES_DIR)
    return [f for f in os.listdir(TEMPLATES_DIR) if f.endswith(".docx") and not f.startswith("~$")]

template_files = get_templates()

if not template_files:
    st.error("لم يتم العثور على أي ملفات Word داخل مجلد templates!")
else:
    selected_template = st.selectbox("اختر النموذج المطلوب معالجته:", template_files)

    st.markdown("---")
    st.subheader("إدخال البيانات المتغيرة")

    col1, col2 = st.columns(2)

    with col1:
        kyc_subject = st.text_input("KYC Subject / Main Dealer")
        client_name_english = st.text_input("Client Name (ENGLISH)")
        client_id = st.text_input("Client ID")
        mcc = st.text_input("MCC")

    with col2:
        region = st.text_input("Region")
        client_name_arabic = st.text_input("Client Name (ARABIC)")
        nature_of_activity = st.text_input("Nature of Activity")

    # قاموس البيانات المستبدلة
    data = {
        "KYC_SUBJECT": kyc_subject,
        "REGION": region,
        "CLIENT_NAME_HEADER": client_name_english,
        "CLIENT_NAME_CELL": client_name_arabic,
        "CLIENT_ID": client_id,
        "NATURE_OF_ACTIVITY": nature_of_activity,
        "MCC": mcc
    }

    st.markdown("---")

    def replace_placeholders(doc, data):
        replacements = {f"{{{{{key}}}}}" : val for key, val in data.items()}
        
        # استبدال في الفقرات العادية
        for p in doc.paragraphs:
            for placeholder, value in replacements.items():
                if placeholder in p.text:
                    p.text = p.text.replace(placeholder, value)

        # استبدال داخل الجداول
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for placeholder, value in replacements.items():
                            if placeholder in p.text:
                                p.text = p.text.replace(placeholder, value)

    def convert_to_pdf_libreoffice(docx_path, output_dir):
        cmd = f"libreoffice --headless --convert-to pdf {docx_path} --outdir {output_dir}"
        subprocess.run(cmd, shell=True, check=True)

    if st.button("معالجة وتجهيز المستند (PDF)", type="primary"):
        template_path = os.path.join(TEMPLATES_DIR, selected_template)
        try:
            doc = Document(template_path)
            replace_placeholders(doc, data)

            temp_docx = "temp_processed.docx"
            doc.save(temp_docx)

            try:
                # تحويل المستند إلى PDF عبر LibreOffice على السحابة
                convert_to_pdf_libreoffice(temp_docx, ".")
                pdf_filename = "temp_processed.pdf"

                with open(pdf_filename, "rb") as f:
                    pdf_bytes = f.read()

                # تنظيف الملفات المؤقتة
                if os.path.exists(temp_docx): os.remove(temp_docx)
                if os.path.exists(pdf_filename): os.remove(pdf_filename)

                st.success("تم تحديث المستند وتحويله إلى PDF بنجاح!")
                st.download_button(
                    label="تنزيل المستند بصيغة PDF",
                    data=pdf_bytes,
                    file_name=f"Updated_{os.path.splitext(selected_template)[0]}.pdf",
                    mime="application/pdf"
                )
            except Exception as pdf_err:
                st.warning("تعذر تحويل الملف إلى PDF تلقائياً، يمكنك تنزيله بصيغة DOCX:")
                bio = io.BytesIO()
                doc.save(bio)
                bio.seek(0)
                st.download_button(
                    label="تنزيل المستند بصيغة DOCX",
                    data=bio,
                    file_name=f"Updated_{selected_template}",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )

        except Exception as e:
            st.error(f"حدث خطأ أثناء معالجة المستند: {e}")